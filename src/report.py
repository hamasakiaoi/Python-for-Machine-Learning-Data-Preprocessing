"""Generate the formal report from executed pipeline artefacts and metadata."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def add_table(doc: Document, df: pd.DataFrame, max_rows: int | None = None) -> None:
    shown = df.head(max_rows) if max_rows else df
    table = doc.add_table(rows=1, cols=len(shown.columns)); table.style = "Light Shading Accent 1"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, name in zip(table.rows[0].cells, shown.columns): cell.text = str(name)
    for _, row in shown.iterrows():
        cells = table.add_row().cells
        for cell, value in zip(cells, row): cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for run in p.runs: run.font.size = Pt(8)

def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)

def add_page_number(paragraph) -> None:
    """Insert a Word PAGE field so pagination updates in Word/LibreOffice."""
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)

def generate_report(root: Path = ROOT) -> Path:
    tables, figures = root / "outputs/tables", root / "outputs/figures"
    meta = json.loads((tables / "run_metadata.json").read_text())
    stages = pd.read_csv(tables / "pipeline_before_after.csv")
    missing = pd.read_csv(tables / "missingness_before.csv")
    scaling = pd.read_csv(tables / "scaling_before_after.csv")
    doc = Document(); section = doc.sections[0]
    section.top_margin = Inches(.75); section.bottom_margin = Inches(.75); section.left_margin = Inches(.8); section.right_margin = Inches(.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
    for name in ["Title", "Heading 1", "Heading 2"]:
        styles[name].font.name = "Aptos Display"; styles[name].font.color.rgb = RGBColor(31, 78, 121)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI Pioneers Internship\n").bold = True
    run = title.add_run("Week 1 — Python for Machine Learning & Data Preprocessing\n\nAdult Census Income: Data Preprocessing and Exploratory Analysis")
    run.font.size = Pt(20); run.font.color.rgb = RGBColor(31,78,121)
    doc.add_paragraph("Student Name: ______________________________", style="Normal").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Machine Learning Internship | 15 September 2026", style="Normal").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    heading(doc, "1. Introduction")
    doc.add_paragraph("This Week 1 project builds a reproducible preprocessing workflow for the UCI Adult Census Income dataset. The work concentrates on the data decisions that precede modelling: inspecting structure, identifying quality issues, treating incomplete values, representing nominal variables, selecting a compact feature set, and standardizing numerical inputs. No predictive model is trained; the final dataset is prepared for a future downstream machine-learning task.")
    heading(doc, "2. Internship Objectives")
    for objective in ["Understand the fundamentals of Machine Learning using Python.", "Learn data loading, cleaning, handling missing values, feature selection, encoding categorical variables, normalization, and exploratory data analysis using Pandas and NumPy.", "Deliverable: Clean and preprocess a sample dataset and document each preprocessing step."]:
        doc.add_paragraph(objective, style="List Bullet")
    heading(doc, "3. Dataset Overview")
    doc.add_paragraph("The Adult Census Income dataset was obtained from the UCI Machine Learning Repository. It was derived from 1994 U.S. Census Bureau data and is commonly used to study whether annual income exceeds $50,000. The raw UCI training and test files were combined only for this preprocessing exercise, yielding 48,842 rows and 15 columns. The source uses '?' as an explicit missing-value marker; the loader converts it to a proper missing value without changing the original files. UCI identifies the dataset as donated by Ronny Kohavi and Barry Becker and links its provenance to Census Bureau data.")
    doc.add_paragraph("Source: https://archive.ics.uci.edu/dataset/2/adult (UCI Machine Learning Repository). The project retains the original files in data/raw and records the source in the README.")
    heading(doc, "4. Initial Data Assessment")
    doc.add_paragraph(f"The raw combined data contains {meta['raw_shape'][0]:,} observations and {meta['raw_shape'][1]} columns. It contains numerical fields (age, sampling weight, education number, capital gain/loss, and weekly hours), nominal demographic/work fields, and the binary income label. {meta['cleaning']['duplicates_removed']} exact duplicate rows were found. The sampling-weight field (fnlwgt) was treated as a survey-design weight rather than a person-level predictor.")
    heading(doc, "5. Exploratory Data Analysis")
    doc.add_paragraph("EDA was deliberately limited to plots that support a preprocessing decision. The missingness chart identifies the fields requiring treatment; occupation frequency exposes the categorical distribution; numerical distributions show the large scale differences that motivate standardization; and the correlation heatmap confirms that education and education_num describe the same underlying level.")
    for file, caption in [("missingness_before.png", "Figure 1. Missingness is concentrated in three categorical fields."), ("numeric_distributions.png", "Figure 2. Age and weekly-hours distributions by income group."), ("numeric_correlation.png", "Figure 3. Numerical correlation check used to review redundancy.")]:
        doc.add_picture(str(figures / file), width=Inches(5.9)); p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading(doc, "6. Data Cleaning")
    doc.add_paragraph(f"Problem → Decision → Reason → Result: {meta['cleaning']['duplicates_removed']} exact duplicate records → remove only exact duplicates → repeated records inflate frequency without adding information → {meta['cleaned_shape'][0]:,} unique records remain. The fnlwgt column → remove → it is a census sampling weight, not an individual attribute → 14 columns remain after this step. Extreme capital-gain/loss values were retained because they are valid monetary observations, not data-entry errors.")
    heading(doc, "7. Missing Value Treatment")
    doc.add_paragraph("Missingness was assessed before any transformation. All affected fields are nominal: workclass, occupation, and native_country. Rather than discard thousands of records, the pipeline uses most-frequent-category imputation inside the categorical transformer. This keeps the complete analytic sample and is transparent, though it may slightly increase the most common category. Native country is subsequently excluded during feature selection for a separate low-information/imbalance reason.")
    add_table(doc, missing.rename(columns={"Unnamed: 0":"feature"}).round(2))
    doc.add_paragraph(f"Before treatment: {meta['raw_missing_cells']:,} missing cells. After treatment: {meta['processed_missing_cells']} missing cells.")
    heading(doc, "8. Categorical Encoding")
    doc.add_paragraph("Seven retained categorical fields—workclass, education, marital_status, occupation, relationship, race, and sex—are nominal. They are encoded with one-hot encoding, not label encoding, because numeric labels would falsely imply an order. The encoder is configured to ignore unseen categories when reused downstream. The final transformed feature matrix contains 62 predictors: 4 standardized numerical features and 58 indicator features.")
    heading(doc, "9. Feature Selection")
    doc.add_paragraph("Candidate fields were reviewed using meaning, redundancy, missingness, and leakage/identifier checks. No direct target leakage field or unique personal identifier exists in the UCI data. The following fields were removed:")
    selection = pd.DataFrame([{"removed feature": k, "reason": v} for k, v in meta['feature_selection'].items()])
    add_table(doc, selection)
    doc.add_paragraph("The retained inputs are age, workclass, education, marital_status, occupation, relationship, race, sex, capital_gain, capital_loss, and hours_per_week. The income column remains the label and is not transformed as a predictor.")
    heading(doc, "10. Normalization / Scaling")
    doc.add_paragraph("StandardScaler was applied only to retained numerical predictors. Standardization was selected over min-max scaling because capital gain and capital loss have strongly skewed, valid extreme values; it places features on a comparable zero-mean, unit-standard-deviation scale without forcing every observation into a fixed range. Categorical indicators and the income label were not scaled.")
    add_table(doc, scaling.round(3))
    heading(doc, "11. Final Preprocessed Dataset")
    doc.add_paragraph(f"The processed CSV contains {meta['processed_shape'][0]:,} rows and {meta['processed_shape'][1]} columns: {meta['encoded_feature_count']} transformed predictors plus the income label. It has no remaining missing values. It is saved as data/processed/adult_census_income_processed.csv; a cleaned pre-encoding checkpoint is also saved for inspection.")
    heading(doc, "12. Before/After Summary")
    add_table(doc, stages)
    heading(doc, "13. Key Findings")
    for finding in ["Missing values are confined to three categorical variables, making categorical imputation more appropriate than numerical imputation or wholesale row deletion.", "The combined raw files contain exact duplicates; removing them preserves one observation per unique record.", "Education and education_num are redundant representations of the same concept; retaining only the readable categorical feature avoids duplicate signal.", "Numerical variables have materially different ranges, especially capital gain/loss versus age and hours; standardization makes their representations comparable."]:
        doc.add_paragraph(finding, style="List Bullet")
    heading(doc, "14. Conclusion")
    doc.add_paragraph("The project delivers the required clean dataset and documents every major preprocessing decision with executable code and before/after evidence. It addresses Python-based data loading, cleaning, missing values, categorical encoding, feature selection, numerical normalization, and exploratory analysis while preserving the original source data. The resulting processed dataset is suitable as an input to a subsequent, separately scoped machine-learning exercise.")
    heading(doc, "15. References")
    doc.add_paragraph("UCI Machine Learning Repository. Adult Dataset. https://archive.ics.uci.edu/dataset/2/adult")
    doc.add_paragraph("Pedregosa et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.")
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AI Pioneers Internship | Week 1 Data Preprocessing | Page ")
    add_page_number(footer)
    output = root / "docs/Week_1_ML_Data_Preprocessing_Report.docx"; output.parent.mkdir(exist_ok=True)
    doc.save(output); return output

if __name__ == "__main__": print(generate_report())
